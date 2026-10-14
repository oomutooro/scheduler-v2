import io
import os
from datetime import date

import django

os.environ['DJANGO_SETTINGS_MODULE'] = 'scheduler.settings'
django.setup()

from django.contrib.messages.storage.fallback import FallbackStorage
from django.test import RequestFactory
from openpyxl import load_workbook
from pypdf import PdfReader

from core import views_reports as v
from core.services.season_analytics import weekly_schedule


def req(**q):
    r = RequestFactory().get('/x/', q)
    r.session = {}
    r._messages = FallbackStorage(r)
    return r


html = v.daily_flights_report(req(season='winter', year=2026)).content.decode()
print('page', len(html) // 1024, 'KB | day sections', html.count('class="day-block"'),
      '| removed columns present:', any(w in html for w in ('>Stand<', '>Gate<', '>Counters<')), '| date column present:', '>Date<' in html)

days = weekly_schedule('winter', 2026, date(2026, 10, 25), date(2027, 3, 27))
for d in days:
    print(d['weekday'], len(d['rows']), 'rows,', d['arrivals'], 'arr,', d['departures'], 'dep')
mon = days[0]['rows']
times = [(r['arr_time'] or r['dep_time']) for r in mon]
print('Monday ordered 00:00->23:59:', times == sorted(times))
for r in mon[:3] + [x for x in mon if x['next_day']][:2]:
    print({k: (str(val) if val is not None else '') for k, val in r.items()})

xlsx = v.daily_flights_export(req(season='winter', year=2026), 'xlsx').content
wb = load_workbook(io.BytesIO(xlsx))
ws = wb['Monday']
print('xlsx sheets', wb.sheetnames, '| header', [c.value for c in ws[4]], '| first row', [c.value for c in ws[5]])
pdf = v.daily_flights_export(req(season='winter', year=2026), 'pdf').content
print('pdf pages', len(PdfReader(io.BytesIO(pdf)).pages))
