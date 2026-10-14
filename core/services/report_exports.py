"""Excel and PDF exports for the daily-schedule and season-analytics reports."""

from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

NAVY = '1E3A8A'
HEADER_FILL = PatternFill('solid', fgColor=NAVY)
HEADER_FONT = Font(bold=True, color='FFFFFF')
BAND_FILL = PatternFill('solid', fgColor='DBEAFE')
THIN = Side(style='thin', color='CBD5E1')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAT_SCALE = ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF', end_type='max', end_color='1A56DB')
PIE_LIMIT = 8


def _header(ws: Any, row: int, values: list[Any]) -> None:
    for col, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _widths(ws: Any, widths: list[int]) -> None:
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _title(ws: Any, text: str, sub: str = '') -> None:
    ws['A1'] = text
    ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    if sub:
        ws['A2'] = sub
        ws['A2'].font = Font(italic=True, color='64748B')


def _bar(ws: Any, title: str, cats: Reference, data: Reference, anchor: str, y_title: str, stacked: bool = False, horizontal: bool = False) -> None:
    chart = BarChart()
    chart.type = 'bar' if horizontal else 'col'
    chart.title, chart.y_axis.title = title, y_title
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = 8.5, 18
    if stacked:
        chart.grouping, chart.overlap = 'stacked', 100
    chart.x_axis.delete = chart.y_axis.delete = False
    ws.add_chart(chart, anchor)


def _heat_sheet(ws: Any, first_header: str, rows: list[Any], value_key: str = 'value') -> None:
    _header(ws, 4, [first_header] + [f'{h:02d}' for h in range(24)])
    for r, row in enumerate(rows, 5):
        ws.cell(row=r, column=1, value=row.get('name') or row.get('label')).font = Font(bold=True)
        for c, cell in enumerate(row['cells'], 2):
            ws.cell(row=r, column=c, value=cell[value_key]).alignment = Alignment(horizontal='center')
    if rows:
        ws.conditional_formatting.add(f'B5:Y{4 + len(rows)}', HEAT_SCALE)
    _widths(ws, [16] + [5] * 24)


def _time_text(row: dict[str, Any], key: str) -> str:
    at = row[key]
    if not at:
        return ''
    return at.strftime('%H:%M') + (' (+1)' if key == 'dep_time' and row['next_day'] else '')


DAILY_HEADER = ['Airline', 'Arr Flight', 'Arr Time', 'From', 'Dep Flight', 'Dep Time', 'To']


def _daily_row(row: dict[str, Any]) -> list[str]:
    return [row['airline'], row['arr_flight'], _time_text(row, 'arr_time'), row['origin'],
            row['dep_flight'], _time_text(row, 'dep_time'), row['destination']]


def daily_workbook(days: list[dict[str, Any]], title: str, subtitle: str) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for day in days:
        ws = wb.create_sheet(day['weekday'])
        _title(ws, f"{title}: {day['weekday']}", f"{subtitle} · {day['arrivals']} arrivals, {day['departures']} departures")
        _header(ws, 4, DAILY_HEADER)
        for row in day['rows']:
            ws.append(_daily_row(row))
        ws.freeze_panes = 'A5'
        if day['rows']:
            ws.auto_filter.ref = f"A4:G{4 + len(day['rows'])}"
        _widths(ws, [30, 12, 11, 8, 12, 14, 8])
    return _save(wb)


SUMMARY_HEADER = ['Airline', 'Type', 'Inbound', 'Arr (EAT)', 'From', 'Days', 'Dept (EAT)', 'Outbound', 'To', 'Period']


def _summary_row(row: dict[str, Any]) -> list[str]:
    return [row['airline'], row['type'], row['arr_flight'], _time_text(row, 'arr_time'), row['origin'], row['days'],
            _time_text(row, 'dep_time'), row['dep_flight'], row['destination'], row['period']]


def summary_workbook(tables: dict[str, list[dict[str, Any]]], title: str, subtitle: str) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for key, sheet in (('Pax', 'Passenger'), ('Cargo', 'Cargo')):
        rows = tables[key]
        ws = wb.create_sheet(sheet)
        _title(ws, f'{title}: {sheet} flights', f'{subtitle} · {len(rows)} flights')
        _header(ws, 4, SUMMARY_HEADER)
        for r, row in enumerate(rows, 5):
            ws.append(_summary_row(row))
            for cell in ws[r]:
                cell.border = BORDER
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            if r % 2 == 0:
                for cell in ws[r]:
                    cell.fill = BAND_FILL
        ws.freeze_panes = 'A5'
        if rows:
            ws.auto_filter.ref = f'A4:J{4 + len(rows)}'
        _widths(ws, [28, 8, 11, 11, 8, 26, 12, 11, 8, 26])
    return _save(wb)


def analytics_workbook(d: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    period = f"{d['season'].title()} {d['year']} · {d['start']:%d %b %Y} – {d['end']:%d %b %Y} (EAT)"
    _title(ws, 'Winter Schedule Analytics' if d['season'] == 'winter' else 'Summer Schedule Analytics', period)
    kpis = [
        ('Operating days', d['total_days']), ('Flight requests', d['flight_requests']), ('Total movements', d['total_movements']),
        ('Arrivals', d['arrivals']), ('Departures', d['departures']), ('Average movements / day', d['avg_per_day']),
        ('Busiest weekday (avg/day)', f"{d['busiest_weekday']['name']} ({d['busiest_weekday']['total']})" if d['busiest_weekday'] else '—'),
        ('Quietest weekday (avg/day)', f"{d['quietest_weekday']['name']} ({d['quietest_weekday']['total']})" if d['quietest_weekday'] else '—'),
        ('Busiest date', f"{d['busiest_day']['date']:%a %d %b %Y} ({d['busiest_day']['total']})" if d['busiest_day'] else '—'),
        ('Busiest hours', ', '.join(f"{h['label']} ({h['total']})" for h in d['busiest_hours'])),
        ('Quietest hours', ', '.join(d['quiet_hours'][:8]) or '—'),
        ('Departures with a gate allocated', f"{d['allocation_coverage']}%"),
    ]
    _header(ws, 4, ['Measure', 'Value'])
    for r, (label, value) in enumerate(kpis, 5):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=value)
    _widths(ws, [34, 60])

    wd = wb.create_sheet('Busy Days')
    _title(wd, 'Average movements per weekday', period)
    _header(wd, 4, ['Weekday', 'Dates in period', 'Arrivals (total)', 'Departures (total)', 'Avg arrivals / day', 'Avg departures / day', 'Avg movements / day'])
    for r, row in enumerate(d['weekday_rows'], 5):
        for c, value in enumerate([row['name'], row['dates'], row['arr_total'], row['dep_total'], row['arr'], row['dep'], row['total']], 1):
            wd.cell(row=r, column=c, value=value)
    _widths(wd, [14, 14, 16, 18, 18, 20, 20])
    _bar(wd, 'Average movements per weekday', Reference(wd, min_col=1, min_row=5, max_row=11), Reference(wd, min_col=5, max_col=6, min_row=4, max_row=11), 'A14', 'Movements / day', stacked=True)

    hr = wb.create_sheet('Busy Hours')
    _title(hr, 'Movements by hour of day (whole period)', period)
    _header(hr, 4, ['Hour', 'Arrivals', 'Departures', 'Total', 'Avg / day'])
    for r, row in enumerate(d['hour_rows'], 5):
        for c, value in enumerate([row['label'], row['arr_total'], row['dep_total'], row['total'], row['avg']], 1):
            hr.cell(row=r, column=c, value=value)
    hr.conditional_formatting.add('D5:D28', HEAT_SCALE)
    _widths(hr, [10, 12, 14, 10, 12])
    _bar(hr, 'Movements by hour', Reference(hr, min_col=1, min_row=5, max_row=28), Reference(hr, min_col=2, max_col=3, min_row=4, max_row=28), 'G4', 'Movements', stacked=True)

    hm = wb.create_sheet('Heatmap')
    _title(hm, 'Heatmap: average movements per hour, by weekday', period)
    _heat_sheet(hm, 'Weekday', d['heat_rows'])

    dt = wb.create_sheet('Daily Totals')
    _title(dt, 'Movements per date', period)
    _header(dt, 4, ['Date', 'Weekday', 'Arrivals', 'Departures', 'Total'])
    for r, row in enumerate(d['daily'], 5):
        dt.cell(row=r, column=1, value=row['date']).number_format = 'DD MMM YYYY'
        for c, value in enumerate([row['weekday'], row['arr'], row['dep'], row['total']], 2):
            dt.cell(row=r, column=c, value=value)
    _widths(dt, [14, 12, 10, 12, 8])
    if d['daily']:
        line = LineChart()
        line.title, line.height, line.width = 'Total movements per date', 8.5, 22
        last = 4 + len(d['daily'])
        line.add_data(Reference(dt, min_col=5, min_row=4, max_row=last), titles_from_data=True)
        line.set_categories(Reference(dt, min_col=1, min_row=5, max_row=last))
        line.x_axis.delete = line.y_axis.delete = False
        line.x_axis.number_format = 'dd mmm'
        dt.add_chart(line, 'G4')

    ac = wb.create_sheet('Aircraft')
    _title(ac, 'Aircraft types used', period)
    _header(ac, 4, ['Type', 'Name', 'Size code', 'Movements', 'Share %'])
    for r, row in enumerate(d['aircraft_rows'], 5):
        for c, value in enumerate([row['code'], row['name'], row['size'], row['movements'], row['share']], 1):
            ac.cell(row=r, column=c, value=value)
    _widths(ac, [10, 32, 10, 12, 10])
    if d['aircraft_rows']:
        last = 4 + len(d['aircraft_rows'])
        _bar(ac, 'Movements by aircraft type', Reference(ac, min_col=1, min_row=5, max_row=last), Reference(ac, min_col=4, min_row=4, max_row=last), 'G4', 'Movements')
        pie_last = 4 + min(len(d['aircraft_rows']), PIE_LIMIT)
        pie = PieChart()
        pie.title, pie.height, pie.width = f'Top {min(len(d["aircraft_rows"]), PIE_LIMIT)} aircraft types', 8.5, 12
        pie.add_data(Reference(ac, min_col=4, min_row=4, max_row=pie_last), titles_from_data=True)
        pie.set_categories(Reference(ac, min_col=1, min_row=5, max_row=pie_last))
        ac.add_chart(pie, 'G22')

    al = wb.create_sheet('Airlines')
    _title(al, 'Airlines by movements', period)
    _header(al, 4, ['IATA', 'Airline', 'Movements', 'Share %'])
    for r, row in enumerate(d['airline_rows'], 5):
        for c, value in enumerate([row['code'], row['name'], row['movements'], row['share']], 1):
            al.cell(row=r, column=c, value=value)
    _widths(al, [8, 30, 12, 10])
    if d['airline_rows']:
        _bar(al, 'Movements by airline', Reference(al, min_col=1, min_row=5, max_row=4 + len(d['airline_rows'])), Reference(al, min_col=3, min_row=4, max_row=4 + len(d['airline_rows'])), 'F4', 'Movements')

    bd = wb.create_sheet('Busiest Dates')
    _title(bd, 'Ten busiest and five quietest dates', period)
    _header(bd, 4, ['Rank', 'Date', 'Weekday', 'Arrivals', 'Departures', 'Total'])
    for r, row in enumerate(d['busiest_dates'], 5):
        for c, value in enumerate([r - 4, row['date'], row['weekday'], row['arr'], row['dep'], row['total']], 1):
            cell = bd.cell(row=r, column=c, value=value)
            if isinstance(value, date):
                cell.number_format = 'DD MMM YYYY'
    start = 6 + len(d['busiest_dates'])
    _header(bd, start, ['Quiet', 'Date', 'Weekday', 'Arrivals', 'Departures', 'Total'])
    for r, row in enumerate(d['quietest_dates'], start + 1):
        for c, value in enumerate([r - start, row['date'], row['weekday'], row['arr'], row['dep'], row['total']], 1):
            cell = bd.cell(row=r, column=c, value=value)
            if isinstance(value, date):
                cell.number_format = 'DD MMM YYYY'
    _widths(bd, [8, 14, 12, 10, 12, 8])

    for key, name, first in (('gates', 'Gate Usage', 'Gate'), ('counters', 'Counter Usage', 'Counter'), ('stands', 'Stand Usage', 'Stand')):
        sheet = wb.create_sheet(name)
        _title(sheet, f'{name}: % of period dates each {first.lower()} is occupied, by hour', period)
        usage = d['resources'][key]
        rows = [{'label': r['label'], 'cells': [{'value': c['percent']} for c in r['cells']]} for r in usage['rows']]
        _heat_sheet(sheet, first, rows)
        _header(sheet, 6 + len(rows), [first, 'Allocations', 'Busiest hour'])
        for r, row in enumerate(usage['rows'], 7 + len(rows)):
            for c, value in enumerate([row['label'], row['allocations'], row['busiest_hour']], 1):
                sheet.cell(row=r, column=c, value=value)
        if usage['rows']:
            first_row, last_row = 7 + len(rows), 6 + 2 * len(rows)
            _bar(sheet, f'{name}: allocations', Reference(sheet, min_col=1, min_row=first_row, max_row=last_row), Reference(sheet, min_col=2, min_row=first_row - 1, max_row=last_row), f'A{last_row + 3}', 'Allocations')
    return _save(wb)


def _save(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ─── PDF ─────────────────────────────────────────────────────────────────────

_styles = getSampleStyleSheet()
H1 = ParagraphStyle('H1', parent=_styles['Heading1'], textColor=colors.HexColor('#' + NAVY), fontSize=18, spaceAfter=4)
H2 = ParagraphStyle('H2', parent=_styles['Heading2'], textColor=colors.HexColor('#1E40AF'), fontSize=13, spaceBefore=10, spaceAfter=4)
SMALL = ParagraphStyle('Small', parent=_styles['Normal'], fontSize=8.5, textColor=colors.HexColor('#475569'))
BODY = ParagraphStyle('Body', parent=_styles['Normal'], fontSize=9.5, leading=13)
PAGE = landscape(A4)
WIDTH = PAGE[0] - 3 * cm


def _blue(intensity: float) -> colors.Color:
    t = max(0.0, min(1.0, intensity))
    return colors.Color(1 - t * (1 - 26 / 255), 1 - t * (1 - 86 / 255), 1 - t * (1 - 219 / 255))


def _table(rows: list[list[Any]], widths: list[float] | None = None, repeat: int = 1, font: float = 8) -> Table:
    table = Table(rows, colWidths=widths, repeatRows=repeat)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#' + NAVY)), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), font),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    return table


def _heat_table(labels: list[str], matrix: list[list[float]], peak: float, label_width: float = 2.6 * cm) -> Table:
    rows = [[''] + [f'{h:02d}' for h in range(24)]] + [[label] + [f'{v:g}' if v else '' for v in row] for label, row in zip(labels, matrix)]
    cell = (WIDTH - label_width) / 24
    table = Table(rows, colWidths=[label_width] + [cell] * 24)
    style = [
        ('FONTSIZE', (0, 0), (-1, -1), 6.5), ('ALIGN', (1, 0), (-1, -1), 'CENTER'), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTNAME', (0, 1), (0, -1), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.2, colors.HexColor('#E2E8F0')),
    ]
    for r, row in enumerate(matrix, 1):
        for c, value in enumerate(row, 1):
            intensity = value / peak if peak else 0
            style.append(('BACKGROUND', (c, r), (c, r), _blue(0.08 + 0.92 * intensity) if value else colors.white))
            if intensity > 0.55:
                style.append(('TEXTCOLOR', (c, r), (c, r), colors.white))
    table.setStyle(TableStyle(style))
    return table


def _bar_drawing(labels: list[str], series: list[list[float]], names: list[str], height: float = 190, width: float | None = None) -> Drawing:
    width = width or WIDTH
    drawing = Drawing(width, height)
    chart = VerticalBarChart()
    chart.x, chart.y, chart.width, chart.height = 40, 38, width - 60, height - 64
    chart.data = series
    chart.categoryAxis.categoryNames = labels
    chart.categoryAxis.labels.fontSize = 7
    chart.categoryAxis.labels.angle = 0 if len(labels) <= 12 else 90
    chart.categoryAxis.labels.dy = -2 if len(labels) <= 12 else -12
    chart.valueAxis.labels.fontSize = 7
    chart.valueAxis.valueMin = 0
    chart.groupSpacing = 4
    palette = [colors.HexColor('#1A56DB'), colors.HexColor('#F59E0B')]
    for i in range(len(series)):
        chart.bars[i].fillColor = palette[i % 2]
    drawing.add(chart)
    for i, name in enumerate(names):
        drawing.add(String(chart.x + 5 + i * 90, height - 12, f'■ {name}', fontSize=8, fillColor=palette[i % 2]))
    return drawing


def _pie_drawing(labels: list[str], values: list[float]) -> Drawing:
    drawing = Drawing(WIDTH / 2, 200)
    pie = Pie()
    pie.x, pie.y, pie.width, pie.height = 20, 30, 140, 140
    pie.data = values
    pie.labels = [f'{label} ({int(v)})' for label, v in zip(labels, values)]
    pie.sideLabels = True
    pie.slices.strokeWidth = 0.5
    pie.slices.fontSize = 7
    palette = ['#1A56DB', '#F59E0B', '#057A55', '#C81E1E', '#7E3AF2', '#0EA5E9', '#64748B', '#EC4899', '#94A3B8']
    for i in range(len(values)):
        pie.slices[i].fillColor = colors.HexColor(palette[i % len(palette)])
    drawing.add(pie)
    return drawing


def _build(story: list[Any], footer: str) -> bytes:
    buffer = BytesIO()

    def decorate(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont('Helvetica', 7.5)
        canvas.setFillColor(colors.HexColor('#64748B'))
        canvas.drawString(1.5 * cm, 0.8 * cm, footer)
        canvas.drawRightString(PAGE[0] - 1.5 * cm, 0.8 * cm, f'Page {doc.page}')
        canvas.restoreState()

    SimpleDocTemplate(buffer, pagesize=PAGE, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.4 * cm, bottomMargin=1.5 * cm).build(
        story, onFirstPage=decorate, onLaterPages=decorate)
    return buffer.getvalue()


def daily_pdf(days: list[dict[str, Any]], title: str, subtitle: str) -> bytes:
    story: list[Any] = [Paragraph(title, H1), Paragraph(subtitle, SMALL), Spacer(1, 6)]
    widths = [7.4 * cm, 3.2 * cm, 2.6 * cm, 2.2 * cm, 3.2 * cm, 3.0 * cm, 2.2 * cm]
    for index, day in enumerate(days):
        if index:
            story.append(PageBreak())
        story.append(Paragraph(
            f"{day['weekday']} <font size=9 color='#64748B'>— {day['arrivals']} arrivals, {day['departures']} departures</font>", H2))
        if day['rows']:
            story.append(_table([DAILY_HEADER] + [_daily_row(r) for r in day['rows']], widths, font=8))
        else:
            story.append(Paragraph('No flights on this day.', BODY))
    return _build(story, f'{title} · {subtitle}')


def analytics_pdf(d: dict[str, Any]) -> bytes:
    period = f"{d['season'].title()} {d['year']} · {d['start']:%d %b %Y} – {d['end']:%d %b %Y} · times in EAT"
    story: list[Any] = [Paragraph(f"{d['season'].title()} {d['year']} Schedule Analytics", H1), Paragraph(period, SMALL), Spacer(1, 8)]

    kpi = [['Measure', 'Value'], ['Operating days', d['total_days']], ['Flight requests', d['flight_requests']],
           ['Total movements', d['total_movements']], ['Arrivals / departures', f"{d['arrivals']} / {d['departures']}"],
           ['Average movements per day', d['avg_per_day']]]
    if d['busiest_weekday']:
        kpi += [['Busiest weekday', f"{d['busiest_weekday']['name']} ({d['busiest_weekday']['total']} / day)"],
                ['Quietest weekday', f"{d['quietest_weekday']['name']} ({d['quietest_weekday']['total']} / day)"],
                ['Busiest date', f"{d['busiest_day']['date']:%a %d %b %Y} ({d['busiest_day']['total']})"],
                ['Busiest hours', ', '.join(f"{h['label']} ({h['total']})" for h in d['busiest_hours'])],
                ['Quietest hours', ', '.join(d['quiet_hours'][:8])]]
    story += [_table(kpi, [7 * cm, 12 * cm], font=9), Spacer(1, 6)]

    story += [KeepTogether([Paragraph('How busy is each day of the week', H2),
                            _bar_drawing([r['name'][:3] for r in d['weekday_rows']], [[r['arr'] for r in d['weekday_rows']], [r['dep'] for r in d['weekday_rows']]], ['Arrivals / day', 'Departures / day'])])]
    story += [KeepTogether([Paragraph('Which time of day is busiest', H2),
                            _bar_drawing([f"{r['hour']:02d}" for r in d['hour_rows']], [[r['arr_total'] for r in d['hour_rows']], [r['dep_total'] for r in d['hour_rows']]], ['Arrivals', 'Departures'])])]

    story += [KeepTogether([Paragraph('Heatmap: average movements per hour by weekday', H2),
                            _heat_table([r['name'][:9] for r in d['heat_rows']], [[c['value'] for c in r['cells']] for r in d['heat_rows']], d['heat_peak']),
                            Paragraph('Darker cells = more movements. Values are the average number of movements in that hour on that weekday.', SMALL)])]

    aircraft = d['aircraft_rows']
    top = aircraft[:PIE_LIMIT]
    if not aircraft:
        story.append(Paragraph('Aircraft types used', H2))
    else:
        chart_row = Table([[_bar_drawing([r['code'] for r in top], [[r['movements'] for r in top]], ['Movements'], 190, WIDTH / 2 - 10), _pie_drawing([r['code'] for r in top], [r['movements'] for r in top])]], colWidths=[WIDTH / 2, WIDTH / 2])
        story.append(KeepTogether([Paragraph('Aircraft types used', H2), chart_row]))
        story.append(_table([['Type', 'Name', 'Size', 'Movements', 'Share %']] + [[r['code'], r['name'], r['size'], r['movements'], r['share']] for r in aircraft], [2.5 * cm, 9 * cm, 2 * cm, 3 * cm, 2.5 * cm]))
    story.append(PageBreak())

    story += [Paragraph('Airlines by movements', H2),
              _table([['IATA', 'Airline', 'Movements', 'Share %']] + [[r['code'], r['name'], r['movements'], r['share']] for r in d['airline_rows']], [2 * cm, 10 * cm, 3 * cm, 3 * cm])]
    story += [Paragraph('Busiest and quietest dates', H2),
              _table([['', 'Date', 'Weekday', 'Arr', 'Dep', 'Total']]
                     + [['Busy', f"{r['date']:%d %b %Y}", r['weekday'], r['arr'], r['dep'], r['total']] for r in d['busiest_dates']]
                     + [['Quiet', f"{r['date']:%d %b %Y}", r['weekday'], r['arr'], r['dep'], r['total']] for r in d['quietest_dates']],
                     [2 * cm, 3.5 * cm, 3.5 * cm, 2 * cm, 2 * cm, 2 * cm])]
    story.append(PageBreak())

    for key, name in (('gates', 'Gate usage'), ('counters', 'Check-in counter usage'), ('stands', 'Stand usage')):
        usage = d['resources'][key]
        rows = usage['rows']
        story.append(Paragraph(f'{name}: % of dates occupied, by hour', H2))
        if any(r['allocations'] for r in rows):
            story.append(_heat_table([r['label'] for r in rows], [[c['percent'] for c in r['cells']] for r in rows],
                                     max(c['percent'] for r in rows for c in r['cells']) or 1))
            story.append(Spacer(1, 4))
            story.append(_table([['Resource', 'Allocations', 'Busiest hour']] + [[r['label'], r['allocations'], r['busiest_hour']] for r in rows if r['allocations']], [4 * cm, 3 * cm, 3 * cm]))
        else:
            story.append(Paragraph('No allocations have been saved for this period yet.', BODY))
    return _build(story, f"Schedule analytics · {period}")
